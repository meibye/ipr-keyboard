"""Shared python-docx helpers for building the Danish manuals.

Keeps build_manuals.py readable: page setup, styles, headings, tables,
code blocks, figures with numbered captions, TOC field and page footer.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FIGURES = Path(__file__).resolve().parents[1] / "figures"

INK = RGBColor(0x1A, 0x20, 0x2C)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
ACCENT = RGBColor(0x2B, 0x6C, 0xB0)
MUTED = RGBColor(0x60, 0x6C, 0x80)
DANGER = RGBColor(0xB0, 0x2A, 0x2A)

BODY_FONT = "Segoe UI"
MONO_FONT = "Consolas"


def _shade(el, hexcolor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    el.append(shd)


def _borders(el, hexcolor: str, sz: int = 4, sides=("top", "left", "bottom", "right")) -> None:
    pbdr = OxmlElement("w:pBdr")
    for side in sides:
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "6")
        b.set(qn("w:color"), hexcolor)
        pbdr.append(b)
    el.append(pbdr)


class Manual:
    def __init__(self, title: str, subtitle: str, doc_id: str, version: str, date: str):
        self.doc = Document()
        self.fig_no = 0
        self.tbl_no = 0
        self.title_text = title
        self._setup_page()
        self._setup_styles()
        self._footer(f"{doc_id} · {title} · version {version}")
        self._cover(title, subtitle, doc_id, version, date)

    # -- setup -------------------------------------------------------------
    def _setup_page(self) -> None:
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.4)
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.0)

    def _setup_styles(self) -> None:
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = BODY_FONT
        n.font.size = Pt(10.5)
        n.font.color.rgb = INK
        n._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        n.paragraph_format.space_after = Pt(7)
        n.paragraph_format.line_spacing = 1.12

        for name, size, color, before, after in (
            ("Heading 1", 19, BLUE, 20, 8),
            ("Heading 2", 14.5, BLUE, 15, 6),
            ("Heading 3", 12, ACCENT, 12, 4),
        ):
            h = st[name]
            h.font.name = BODY_FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = color
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True

        for lname in ("List Bullet", "List Number"):
            ls = st[lname]
            ls.font.name = BODY_FONT
            ls.font.size = Pt(10.5)
            ls.paragraph_format.space_after = Pt(3)

    def _footer(self, text: str) -> None:
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text + "   ·   Side ")
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        r.font.name = BODY_FONT
        self._field(p, "PAGE")
        r2 = p.add_run(" af ")
        r2.font.size = Pt(8)
        r2.font.color.rgb = MUTED
        r2.font.name = BODY_FONT
        self._field(p, "NUMPAGES")

    @staticmethod
    def _field(paragraph, instr: str) -> None:
        r = paragraph.add_run()
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        r.font.name = BODY_FONT
        for tag, attr, txt in (
            ("w:fldChar", ("w:fldCharType", "begin"), None),
            ("w:instrText", ("xml:space", "preserve"), f" {instr} "),
            ("w:fldChar", ("w:fldCharType", "end"), None),
        ):
            el = OxmlElement(tag)
            el.set(qn(attr[0]), attr[1])
            if txt is not None:
                el.text = txt
            r._r.append(el)

    # -- cover / toc -------------------------------------------------------
    def _cover(self, title, subtitle, doc_id, version, date) -> None:
        for _ in range(4):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.size = Pt(30)
        r.font.bold = True
        r.font.color.rgb = BLUE
        r.font.name = BODY_FONT

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        r.font.size = Pt(13)
        r.font.color.rgb = MUTED
        r.font.name = BODY_FONT

        self.doc.add_paragraph()
        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in (
            ("Dokument-id", doc_id),
            ("Version", version),
            ("Dato", date),
            ("Produkt", "IPR Pen Bridge — Raspberry Pi Zero 2 W"),
            ("Sprog", "Dansk"),
        ):
            row = t.add_row().cells
            rk = row[0].paragraphs[0].add_run(k)
            rk.bold = True
            rk.font.size = Pt(10)
            rk.font.name = BODY_FONT
            rv = row[1].paragraphs[0].add_run(v)
            rv.font.size = Pt(10)
            rv.font.name = BODY_FONT
        for row in t.rows:
            row.cells[0].width = Cm(4.5)
            row.cells[1].width = Cm(9.5)

        self.page_break()

    def toc(self) -> None:
        self.h1("Indhold")
        p = self.doc.add_paragraph()
        r = p.add_run()
        r.font.name = BODY_FONT
        for tag, attr, txt in (
            ("w:fldChar", ("w:fldCharType", "begin"), None),
            ("w:instrText", ("xml:space", "preserve"), r' TOC \o "1-2" \h \z \u '),
            ("w:fldChar", ("w:fldCharType", "separate"), None),
        ):
            el = OxmlElement(tag)
            el.set(qn(attr[0]), attr[1])
            if txt is not None:
                el.text = txt
            r._r.append(el)
        r2 = p.add_run("Højreklik her i Word og vælg “Opdatér felt” (eller tryk Ctrl+A og F9) "
                       "for at udfylde indholdsfortegnelsen.")
        r2.italic = True
        r2.font.size = Pt(9)
        r2.font.color.rgb = MUTED
        r2.font.name = BODY_FONT
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "end")
        r2._r.append(el)
        self.page_break()

    # -- content -----------------------------------------------------------
    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def h1(self, text: str, new_page: bool = False):
        if new_page:
            self.page_break()
        return self.doc.add_heading(text, level=1)

    def h2(self, text: str):
        return self.doc.add_heading(text, level=2)

    def h3(self, text: str):
        return self.doc.add_heading(text, level=3)

    def p(self, text: str = "", bold: bool = False, italic: bool = False,
          size: float = 10.5, color=None):
        par = self.doc.add_paragraph()
        r = par.add_run(text)
        r.bold, r.italic = bold, italic
        r.font.size = Pt(size)
        r.font.name = BODY_FONT
        if color is not None:
            r.font.color.rgb = color
        return par

    def rich(self, parts):
        """parts: iterable of (text, style) where style in {'', 'b', 'i', 'c'}."""
        par = self.doc.add_paragraph()
        for text, style in parts:
            r = par.add_run(text)
            r.bold = "b" in style
            r.italic = "i" in style
            if "c" in style:
                r.font.name = MONO_FONT
                r.font.size = Pt(9.5)
            else:
                r.font.name = BODY_FONT
                r.font.size = Pt(10.5)
        return par

    def bullets(self, items, numbered: bool = False):
        style = "List Number" if numbered else "List Bullet"
        for it in items:
            par = self.doc.add_paragraph(style=style)
            if isinstance(it, tuple):
                lead, rest = it
                r = par.add_run(lead)
                r.bold = True
                r.font.name = BODY_FONT
                r.font.size = Pt(10.5)
                r2 = par.add_run(rest)
                r2.font.name = BODY_FONT
                r2.font.size = Pt(10.5)
            else:
                r = par.add_run(it)
                r.font.name = BODY_FONT
                r.font.size = Pt(10.5)

    def code(self, text: str):
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.left_indent = Cm(0.4)
        pf.space_before = Pt(4)
        pf.space_after = Pt(8)
        pf.line_spacing = 1.0
        _shade(par._p.get_or_add_pPr(), "F4F6F8")
        _borders(par._p.get_or_add_pPr(), "DCE1E8")
        for i, line in enumerate(text.strip("\n").split("\n")):
            r = par.add_run(line)
            r.font.name = MONO_FONT
            r.font.size = Pt(9)
            if i < len(text.strip("\n").split("\n")) - 1:
                r.add_break()
        return par

    def note(self, text: str, kind: str = "info"):
        colors = {"info": ("EBF4FF", "2B6CB0", "Bemærk"),
                  "warn": ("FFF7E6", "9A6A00", "Vigtigt"),
                  "danger": ("FEF0F0", "B02A2A", "Advarsel"),
                  "tip": ("E9F7EF", "277646", "Tip")}
        fill, border, label = colors[kind]
        par = self.doc.add_paragraph()
        pf = par.paragraph_format
        pf.left_indent = Cm(0.2)
        pf.space_before = Pt(6)
        pf.space_after = Pt(10)
        _shade(par._p.get_or_add_pPr(), fill)
        _borders(par._p.get_or_add_pPr(), border, sz=6)
        r = par.add_run(label + ": ")
        r.bold = True
        r.font.name = BODY_FONT
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(border)
        r2 = par.add_run(text)
        r2.font.name = BODY_FONT
        r2.font.size = Pt(10)
        return par

    def table(self, headers, rows, widths=None, caption: str | None = None,
              mono_cols=()):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0].cells
        for i, htext in enumerate(headers):
            hdr[i].paragraphs[0].paragraph_format.space_after = Pt(2)
            r = hdr[i].paragraphs[0].add_run(htext)
            r.bold = True
            r.font.size = Pt(9.5)
            r.font.name = BODY_FONT
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(hdr[i]._tc.get_or_add_tcPr(), "2B6CB0")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                par = cells[i].paragraphs[0]
                par.paragraph_format.space_after = Pt(2)
                first = True
                for seg in str(val).split("\n"):
                    if not first:
                        par.add_run().add_break()
                    r = par.add_run(seg)
                    r.font.size = Pt(9.5)
                    r.font.name = MONO_FONT if i in mono_cols else BODY_FONT
                    if i in mono_cols:
                        r.font.size = Pt(9)
                    first = False
                if ri % 2 == 1:
                    _shade(cells[i]._tc.get_or_add_tcPr(), "F7FAFC")
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        if caption:
            self.tbl_no += 1
            self._caption(f"Tabel {self.tbl_no}. {caption}")
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def _caption(self, text: str) -> None:
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(3)
        par.paragraph_format.space_after = Pt(12)
        r = par.add_run(text)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        r.font.name = BODY_FONT

    def figure(self, filename: str, caption: str, width_cm: float = 15.6) -> None:
        path = FIGURES / filename
        if not path.exists():
            raise FileNotFoundError(path)
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.space_before = Pt(8)
        par.paragraph_format.space_after = Pt(2)
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self.fig_no += 1
        self._caption(f"Figur {self.fig_no}. {caption}")

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        print(f"  wrote {path.name}  ({self.fig_no} figurer, {self.tbl_no} tabeller)")
